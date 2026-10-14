"""산출근거서 엑셀: 시트 구성과 실별 산출서 합계가 AI 산출과 같은지 확인한다."""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "site/assets/산출근거서_대한변협신회관_바닥마감.xlsx"
V6 = json.loads((ROOT / "runs/v6/output.json").read_text())


def build():
    subprocess.run([sys.executable, str(ROOT / "harness/make_xlsx.py")], check=True, capture_output=True)
    return load_workbook(XLSX)


def test_sheets():
    wb = build()
    assert wb.sheetnames == ["공종별집계", "실별 산출서", "산출근거서", "검토 필요", "실험 기록"]


def test_room_sheet_matches_ai_totals():
    ws = build()["실별 산출서"]
    a_values, floor_rows, bb_rows = {}, 0, 0
    total_floor = total_bb = 0.0
    for row in ws.iter_rows(min_row=1):
        part, qty = row[1].value, row[6].value
        if row[0].value in ("A", "B"):
            a_values[row[6].coordinate] = qty
        if part in ("바닥", "걸레받이"):
            assert isinstance(qty, str) and qty.startswith("=G"), "물량은 변수 셀 참조 수식이어야 함"
            ref = qty[1:].split("-")[0]
            val = a_values[ref]
            if part == "바닥":
                floor_rows += 1; total_floor += val
            else:
                bb_rows += 1; total_bb += val
    assert floor_rows == bb_rows == len(V6["rooms"]) == 36
    assert round(total_floor, 2) == V6["totals"]["floor_total"]
    assert round(total_bb, 2) == V6["totals"]["baseboard_m"]
